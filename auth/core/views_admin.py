import logging
import secrets
import string
from functools import wraps

from django.contrib import messages
from django.contrib.auth import update_session_auth_hash
from django.contrib.auth.decorators import login_required
from django.db.models import Count
from django.shortcuts import render, redirect, get_object_or_404

from core.models import Classe, Utilisateur

logger = logging.getLogger('auth')


def admin_requis(view_func):
    """Accès réservé aux administrateurs."""
    @wraps(view_func)
    @login_required
    def wrapper(request, *args, **kwargs):
        if request.user.role != 'admin':
            return redirect('connexion')
        return view_func(request, *args, **kwargs)
    return wrapper


def generer_mot_de_passe(longueur=10):
    alphabet = string.ascii_letters + string.digits
    return ''.join(secrets.choice(alphabet) for _ in range(longueur))


# ── Page principale : import de classe ──────────────────────────

@admin_requis
def gestion(request):
    resultats = None

    if request.method == 'POST':
        resultats = _traiter_import(request)

    classes = Classe.objects.annotate(
        nb_etudiants=Count('utilisateurs')
    ).order_by('-annee', 'nom')

    return render(request, 'gestion.html', {
        'active_page': 'gestion',
        'classes': classes,
        'resultats': resultats,
    })


def _traiter_import(request):
    fichier = request.FILES.get('fichier')
    annee = request.POST.get('annee', '').strip()
    generer_mdp = 'generer_mdp' in request.POST
    debut = request.POST.get('debut_autorisation') or None
    fin = request.POST.get('fin_autorisation') or None

    if not fichier or not annee:
        messages.error(request, "Veuillez fournir un fichier et une année.")
        return None

    if not fichier.name.endswith('.xlsx'):
        messages.error(request, "Le fichier doit être au format .xlsx")
        return None

    try:
        import openpyxl
        wb = openpyxl.load_workbook(fichier, read_only=True)
    except Exception:
        messages.error(request, "Impossible de lire le fichier Excel.")
        return None

    resultats = None

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))

        if len(rows) < 2:
            messages.warning(request, f"Feuille « {sheet_name} » ignorée (vide).")
            continue

        # Trouver les colonnes par en-tête
        header = [str(c).strip().lower() if c else '' for c in rows[0]]
        col_nom = None
        col_email = None
        for i, h in enumerate(header):
            if 'nom' in h and ('prénom' in h or 'prenom' in h):
                col_nom = i
            elif 'email' in h or 'mail' in h:
                col_email = i

        if col_nom is None or col_email is None:
            messages.error(request, f"Feuille « {sheet_name} » : colonnes NOM Prénom et/ou Email introuvables.")
            continue

        # Créer ou récupérer la classe
        classe, _ = Classe.objects.get_or_create(nom=sheet_name, annee=annee)
        if debut:
            classe.debut_autorisation = debut
        if fin:
            classe.fin_autorisation = fin
        if debut or fin:
            classe.save()

        comptes = []
        crees = 0
        existants = 0

        for row in rows[1:]:
            nom_complet = row[col_nom] if col_nom < len(row) else None
            email = row[col_email] if col_email < len(row) else None

            if not nom_complet or not email:
                continue

            nom_complet = str(nom_complet).strip()
            email = str(email).strip().lower()
            mot_de_passe = generer_mot_de_passe() if generer_mdp else None

            try:
                user = Utilisateur.objects.get(email=email)
                user.classes.add(classe)
                existants += 1
                comptes.append({
                    'nom': nom_complet,
                    'email': email,
                    'mot_de_passe': '(inchangé)',
                    'nouveau': False,
                })
            except Utilisateur.DoesNotExist:
                user = Utilisateur.objects.create_user(
                    email=email,
                    password=mot_de_passe,
                    nom_complet=nom_complet,
                    role='etudiant',
                )
                user.classes.add(classe)
                crees += 1
                comptes.append({
                    'nom': nom_complet,
                    'email': email,
                    'mot_de_passe': mot_de_passe or '(non généré)',
                    'nouveau': True,
                })

        resultats = {
            'classe': f"{sheet_name} ({annee})",
            'crees': crees,
            'existants': existants,
            'comptes': comptes,
        }

        logger.info(
            f"Import classe {sheet_name}: {crees} créés, {existants} existants",
            extra={'user': request.user.id, 'classe': sheet_name, 'ip': '-',
                   'method': 'POST', 'path': '/admin/', 'status': 200}
        )

    wb.close()

    if resultats:
        messages.success(request, "Import terminé avec succès.")

    return resultats


# ── Détail d'une classe ─────────────────────────────────────────

@admin_requis
def detail_classe(request, classe_id):
    classe = get_object_or_404(Classe, id=classe_id)
    etudiants = classe.utilisateurs.filter(role='etudiant').order_by('nom_complet')
    return render(request, 'detail_classe.html', {
        'active_page': 'gestion',
        'classe': classe,
        'etudiants': etudiants,
    })


# ── Profil utilisateur (changement de mot de passe) ────────────

@login_required
def profil(request):
    if request.method == 'POST':
        ancien = request.POST.get('ancien_mdp', '')
        nouveau = request.POST.get('nouveau_mdp', '')
        confirmation = request.POST.get('confirmation_mdp', '')

        if not request.user.check_password(ancien):
            messages.error(request, "Mot de passe actuel incorrect.")
        elif len(nouveau) < 8:
            messages.error(request, "Le nouveau mot de passe doit contenir au moins 8 caractères.")
        elif nouveau != confirmation:
            messages.error(request, "Les mots de passe ne correspondent pas.")
        else:
            request.user.set_password(nouveau)
            request.user.save()
            update_session_auth_hash(request, request.user)
            messages.success(request, "Mot de passe modifié avec succès.")

    return render(request, 'profil.html', {'active_page': 'profil'})
