"""Côté admin : import des classes, page d'une classe, archivage puis suppression."""
import logging
import secrets
import string
import tempfile

from flask import Blueprint, abort, flash, redirect, render_template, request, send_file, url_for

from ada import archivage, db
from ada.comptes import admin_requis, creer_utilisateur, utilisateur_courant
from ada.modeles import creer_classe, get_classe, membres_classe, toutes_les_classes

bp = Blueprint('admin', __name__, url_prefix='/admin')
journal = logging.getLogger('ada')


def generer_mot_de_passe(longueur=10):
    alphabet = string.ascii_letters + string.digits
    return ''.join(secrets.choice(alphabet) for _ in range(longueur))


def _classe(classe_id):
    classe = get_classe(classe_id)
    if classe is None:
        abort(404)
    return classe


# ── Page principale : import de classe ──────────────────────────

@bp.route('/', methods=['GET', 'POST'])
@admin_requis
def gestion():
    resultats = None
    if request.method == 'POST':
        resultats = _traiter_import()

    effectifs = dict(db.commun().execute(
        'SELECT classe_id, count(*) FROM inscriptions GROUP BY classe_id'
    ).fetchall())
    classes = [(classe, effectifs.get(classe.id, 0)) for classe in toutes_les_classes()]
    return render_template('gestion.html', active_page='gestion', classes=classes, resultats=resultats)


def _date_ou_none(texte):
    texte = (texte or '').strip()
    return texte or None


def _traiter_import():
    utilisateur = utilisateur_courant()
    fichier = request.files.get('fichier')
    annee = request.form.get('annee', '').strip()
    generer_mdp = 'generer_mdp' in request.form
    debut = _date_ou_none(request.form.get('debut_autorisation'))
    fin = _date_ou_none(request.form.get('fin_autorisation'))

    if not fichier or not fichier.filename or not annee:
        flash("Veuillez fournir un fichier et une année.", 'danger')
        return None

    if not fichier.filename.endswith('.xlsx'):
        flash("Le fichier doit être au format .xlsx", 'danger')
        return None

    try:
        import openpyxl
        wb = openpyxl.load_workbook(fichier.stream, read_only=True)
    except Exception:
        flash("Impossible de lire le fichier Excel.", 'danger')
        return None

    resultats = None
    cx = db.commun()

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))

        if len(rows) < 2:
            flash(f"Feuille « {sheet_name} » ignorée (vide).", 'warning')
            continue

        # Trouver les colonnes par en-tête
        header = [str(c).strip().lower() if c else '' for c in rows[0]]
        col_nom = None
        col_email = None
        for i, h in enumerate(header):
            if 'nom' in h and ('prénom' in h or 'prenom' in h):
                col_nom = i
            elif 'email' in h or 'mail' in h:
                col_email = i

        if col_nom is None or col_email is None:
            flash(f"Feuille « {sheet_name} » : colonnes NOM Prénom et/ou Email introuvables.", 'danger')
            continue

        # Créer ou récupérer la classe (et sa base)
        classe, _cree = creer_classe(sheet_name, annee)
        if debut or fin:
            with db.transaction(cx):
                cx.execute(
                    'UPDATE classes SET debut_autorisation = COALESCE(?, debut_autorisation), '
                    'fin_autorisation = COALESCE(?, fin_autorisation) WHERE id = ?',
                    (debut, fin, classe.id),
                )

        comptes = []
        crees = 0
        existants = 0

        with db.transaction(cx):
            for row in rows[1:]:
                nom_complet = row[col_nom] if col_nom < len(row) else None
                email = row[col_email] if col_email < len(row) else None

                if not nom_complet or not email:
                    continue

                nom_complet = str(nom_complet).strip()
                email = str(email).strip().lower()
                mot_de_passe = generer_mot_de_passe() if generer_mdp else None

                ligne = cx.execute('SELECT id FROM utilisateurs WHERE email = ?', (email,)).fetchone()
                if ligne:
                    utilisateur_id = ligne['id']
                    existants += 1
                    comptes.append({
                        'nom': nom_complet,
                        'email': email,
                        'mot_de_passe': '(inchangé)',
                        'nouveau': False,
                    })
                else:
                    utilisateur_id = creer_utilisateur(cx, email, nom_complet, 'etudiant', mot_de_passe)
                    crees += 1
                    comptes.append({
                        'nom': nom_complet,
                        'email': email,
                        'mot_de_passe': mot_de_passe or '(non généré)',
                        'nouveau': True,
                    })
                cx.execute(
                    'INSERT OR IGNORE INTO inscriptions (utilisateur_id, classe_id) VALUES (?, ?)',
                    (utilisateur_id, classe.id),
                )

        resultats = {
            'classe': f"{sheet_name} ({annee})",
            'crees': crees,
            'existants': existants,
            'comptes': comptes,
        }
        journal.info(
            "Import classe=%s nom=%s crees=%s existants=%s par=%s",
            classe.id, sheet_name, crees, existants, utilisateur.id,
        )

    wb.close()

    if resultats:
        flash("Import terminé avec succès.", 'success')

    return resultats


# ── Détail d'une classe ─────────────────────────────────────────

@bp.get('/classes/<int:classe_id>/')
@admin_requis
def detail_classe(classe_id):
    classe = _classe(classe_id)
    donnees = archivage.donnees_classe(classe)
    return render_template(
        'detail_classe.html',
        active_page='gestion',
        classe=classe,
        etudiants=membres_classe(classe.id, role='etudiant'),
        bilan=archivage.bilan(donnees),
        empreinte=archivage.calculer_empreinte(donnees),
    )


# ── Archivage d'une classe : export, puis suppression ──────────

@bp.get('/classes/<int:classe_id>/archive/')
@admin_requis
def archive_classe(classe_id):
    classe = _classe(classe_id)
    temporaire = tempfile.NamedTemporaryFile(suffix='.zip')
    try:
        resultat = archivage.ecrire_archive(classe, temporaire)
        temporaire.seek(0)
        archivage.verifier_archive(classe, temporaire)
    except Exception:
        temporaire.close()
        raise
    # Nouveau descripteur, à l'offset 0 : après la relecture, seek(0) peut rester dans le tampon
    # de Python sans déplacer le descripteur, que gunicorn envoie tel quel (sendfile).
    envoi = open(temporaire.name, 'rb')
    temporaire.close()
    journal.info(
        "Classe exportée id=%s empreinte=%s fichiers=%s manquants=%s par=%s",
        classe.id, resultat['empreinte'][:8], resultat['fichiers'], resultat['manquants'],
        utilisateur_courant().id,
    )
    return send_file(
        envoi,
        as_attachment=True,
        download_name=archivage.nom_archive(classe, resultat['empreinte']),
        mimetype='application/zip',
    )


@bp.post('/classes/<int:classe_id>/supprimer/')
@admin_requis
def supprimer_classe(classe_id):
    classe = _classe(classe_id)
    if request.form.get('confirmation', '').strip() != classe.nom or 'archive_verifiee' not in request.form:
        flash("Confirmation incorrecte : tapez le nom exact de la classe et cochez la case.", 'danger')
        return redirect(url_for('admin.detail_classe', classe_id=classe.id))
    empreinte = archivage.calculer_empreinte(archivage.donnees_classe(classe))
    if request.form.get('empreinte') != empreinte:
        flash(
            "Les données de la classe ont changé depuis l'affichage de la page : exportez à nouveau l'archive.",
            'danger',
        )
        return redirect(url_for('admin.detail_classe', classe_id=classe.id))
    nom = str(classe)
    resultat = archivage.supprimer_classe(classe)
    flash(
        f"Classe « {nom} » supprimée : {resultat['devoirs']} devoirs, {resultat['fichiers']} fichiers, "
        f"{resultat['etudiants_supprimes']} comptes supprimés, "
        f"{resultat['etudiants_conserves']} conservés (inscrits dans une autre classe).",
        'success',
    )
    journal.info(
        "Classe supprimée id=%s nom=%s empreinte=%s devoirs=%s fichiers=%s comptes_supprimes=%s par=%s",
        classe_id, nom, empreinte[:8], resultat['devoirs'], resultat['fichiers'],
        resultat['etudiants_supprimes'], utilisateur_courant().id,
    )
    return redirect(url_for('admin.gestion'))
